import os
import shutil
import sys
sys.path.insert(0, "/Users/ahmadkhaled/Documents/GitHub/mixed")
from services.word_editor import replace_placeholders
import docx

os.makedirs("/Users/ahmadkhaled/Documents/GitHub/mixed/scratch/color_test", exist_ok=True)
template_path = "/Users/ahmadkhaled/Documents/GitHub/mixed/core_app/templates/template.docx"
output_path = "/Users/ahmadkhaled/Documents/GitHub/mixed/scratch/color_test/test_colored_doc.docx"

test_html_body = """
<h1 style="text-align: center;">خطاب رسمي موحد (Default Navy H1)</h1>
<h2 style="color: #16A34A; text-align: right;">عنوان فرعي باللون الأخضر (Green H2)</h2>
<p style="text-align: right;">
نحيط سيادتكم علماً بأن هذا المستند يتضمن 
<span style="color: #DC2626; font-weight: bold;">نصاً باللون الأحمر العريض</span>، 
وكذلك <span style="color: #2563EB;">نصاً باللون الأزرق</span>، 
وأيضاً <span style="background-color: #FEF08A;">نصاً مميزاً باللون الأصفر الفاقع</span>.
</p>
<ul style="text-align: right;">
    <li>العنصر الأول: نقطة مع <span style="color: #7C3AED;">لون بنفسجي مميز</span></li>
    <li>العنصر الثاني: نقطة باللون الأسود القياسي</li>
</ul>
<ol style="text-align: right;">
    <li>الخطوة الأولى: تأكيد التنسيق الملون</li>
    <li>الخطوة الثانية: <span style="color: #D97706; font-weight: bold;">خطوة ذهبية برتقالية</span></li>
</ol>
"""

placeholders = {
    "{{CODE_NUMBER}}": "1001 HUEED 10-2026",
    "{{BODY_TEXT}}": test_html_body,
    "{{SENDER}}": "أ.د. عميد الكلية",
    "{{RECIPIENT}}": "السادة أعضاء هيئة التدريس",
    "{{SUBJECT}}": "اختبار دعم الألوان والتمييز في المستند",
    "{{HOLDER_NAME}}": "د. مصطفى كمال",
    "{{POSITION}}": "وكيل الكلية لشؤون التعليم والطلاب",
    "{{DATE}}": "2026-10-25"
}

replace_placeholders(
    input_path=template_path,
    output_path=output_path,
    replacements=placeholders
)

print(f"Generated docx at: {output_path}")

doc = docx.Document(output_path)
print(f"Total paragraphs in doc: {len(doc.paragraphs)}")
colored_runs = 0
highlighted_runs = 0
for p_idx, p in enumerate(doc.paragraphs):
    for r_idx, r in enumerate(p.runs):
        c = r.font.color.rgb if (r.font and r.font.color) else None
        hl = r._element.xpath('.//w:highlight/@w:val')
        if c or hl:
            print(f"P{p_idx} R{r_idx} text='{r.text.strip()}' color={c} highlight={hl}")
            if c: colored_runs += 1
            if hl: highlighted_runs += 1

print(f"Verification summary: {colored_runs} colored runs, {highlighted_runs} highlighted runs found!")

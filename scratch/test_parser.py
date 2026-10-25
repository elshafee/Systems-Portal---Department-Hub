from html.parser import HTMLParser
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def parse_color_to_rgb(val):
    if not val:
        return None
    val = val.strip().lower()
    
    # rgb(...)
    rgb_match = re.search(r'rgba?\s*\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)', val)
    if rgb_match:
        try:
            r = min(255, max(0, int(rgb_match.group(1))))
            g = min(255, max(0, int(rgb_match.group(2))))
            b = min(255, max(0, int(rgb_match.group(3))))
            return (r, g, b)
        except:
            pass
            
    # hex #RRGGBB or #RGB
    hex_match = re.search(r'#([0-9a-f]{6}|[0-9a-f]{3})\b', val)
    if hex_match:
        h = hex_match.group(1)
        if len(h) == 3:
            h = ''.join([c*2 for c in h])
        r = int(h[0:2], 16)
        g = int(h[2:4], 16)
        b = int(h[4:6], 16)
        return (r, g, b)

    named_colors = {
        'black': (0, 0, 0),
        'white': (255, 255, 255),
        'red': (220, 38, 38),
        'blue': (37, 99, 235),
        'green': (22, 163, 74),
        'yellow': (234, 179, 8),
        'orange': (234, 88, 12),
        'purple': (124, 58, 237),
        'gray': (107, 114, 128),
        'grey': (107, 114, 128),
        'navy': (31, 78, 121),
        'darkblue': (30, 58, 138),
        'gold': (217, 119, 6),
        'maroon': (128, 0, 0),
        'teal': (13, 148, 136),
    }
    for name, rgb in named_colors.items():
        if f'color:{name}' in val.replace(' ', '') or val == name:
            return rgb
            
    return None

def parse_highlight(val):
    if not val:
        return None
    val = val.strip().lower()
    if 'yellow' in val or '#fef08a' in val or '#ffff00' in val:
        return 'yellow'
    if 'green' in val or '#bbf7d0' in val:
        return 'green'
    if 'blue' in val or 'cyan' in val or '#bae6fd' in val:
        return 'cyan'
    if 'pink' in val or 'magenta' in val or '#fbce8' in val or '#fbcfe8' in val:
        return 'magenta'
    if 'orange' in val or '#fed7aa' in val:
        return 'darkYellow'
    return None

class EnhancedDocxHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.paragraphs = []
        self.current_runs = []
        self.current_align = 'right'
        self.current_heading = 0
        self.list_stack = []
        self.ol_counter_stack = []
        
        self.base_format = {'bold': False, 'italic': False, 'underline': False, 'color': None, 'highlight': None}
        self.format_stack = [self.base_format.copy()]
        self.tag_stack = []

    def _flush_paragraph(self):
        if self.current_runs:
            self.paragraphs.append({
                'runs': self.current_runs,
                'alignment': self.current_align,
                'heading_level': self.current_heading
            })
            self.current_runs = []

    def handle_starttag(self, tag, attrs):
        attrs_dict = dict(attrs)
        style = attrs_dict.get('style', '').lower()
        align_attr = attrs_dict.get('align', '').lower()

        new_align = None
        if 'text-align: center' in style or align_attr == 'center':
            new_align = 'center'
        elif 'text-align: left' in style or align_attr == 'left':
            new_align = 'left'
        elif 'text-align: justify' in style or align_attr == 'justify':
            new_align = 'justify'
        elif 'text-align: right' in style or align_attr == 'right':
            new_align = 'right'

        if tag in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'div']:
            self._flush_paragraph()
            if new_align:
                self.current_align = new_align
            elif not self.list_stack:
                self.current_align = 'right'

            if tag == 'h1':
                self.current_heading = 1
            elif tag == 'h2':
                self.current_heading = 2
            elif tag in ['h3', 'h4', 'h5', 'h6']:
                self.current_heading = 3
            else:
                self.current_heading = 0
        elif tag == 'ul':
            self.list_stack.append('ul')
            self.ol_counter_stack.append(0)
        elif tag == 'ol':
            self.list_stack.append('ol')
            self.ol_counter_stack.append(1)
        elif tag == 'li':
            self._flush_paragraph()
            if new_align:
                self.current_align = new_align
            else:
                self.current_align = 'right'
            self.current_heading = 0

            if self.list_stack and self.list_stack[-1] == 'ol':
                count = self.ol_counter_stack[-1]
                prefix = f'{count}. '
                self.ol_counter_stack[-1] += 1
            else:
                prefix = '• '

            active_fmt = self.format_stack[-1]
            self.current_runs.append({
                'text': prefix,
                'bold': True,
                'italic': False,
                'underline': False,
                'color': active_fmt.get('color'),
                'highlight': None
            })

        # Process formatting overrides
        overrides = {}
        if tag in ['strong', 'b'] or 'font-weight: bold' in style or 'font-weight: 700' in style:
            overrides['bold'] = True
        if tag in ['em', 'i'] or 'font-style: italic' in style:
            overrides['italic'] = True
        if tag in ['u'] or 'text-decoration: underline' in style:
            overrides['underline'] = True

        # Color
        color_val = None
        if 'color:' in style:
            m = re.search(r'(?:^|;)\s*color\s*:\s*([^;]+)', style)
            if m:
                color_val = m.group(1).strip()
        elif tag == 'font' and 'color' in attrs_dict:
            color_val = attrs_dict['color'].strip()

        if color_val:
            rgb = parse_color_to_rgb(color_val)
            if rgb:
                overrides['color'] = rgb

        # Background / Highlight
        bg_val = None
        if 'background-color:' in style:
            m = re.search(r'(?:^|;)\s*background-color\s*:\s*([^;]+)', style)
            if m:
                bg_val = m.group(1).strip()
        elif tag == 'mark':
            bg_val = 'yellow'

        if bg_val:
            hl = parse_highlight(bg_val)
            if hl:
                overrides['highlight'] = hl

        new_fmt = self.format_stack[-1].copy()
        new_fmt.update(overrides)
        self.format_stack.append(new_fmt)
        self.tag_stack.append(tag)

    def handle_endtag(self, tag):
        if tag in self.tag_stack:
            for i in range(len(self.tag_stack) - 1, -1, -1):
                if self.tag_stack[i] == tag:
                    self.tag_stack.pop(i)
                    self.format_stack.pop(i + 1)
                    break

        if tag in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6', 'p', 'div', 'li']:
            self._flush_paragraph()
            self.current_heading = 0
        elif tag in ['ul', 'ol']:
            self._flush_paragraph()
            if self.list_stack:
                self.list_stack.pop()
            if self.ol_counter_stack:
                self.ol_counter_stack.pop()
        elif tag == 'br':
            self._flush_paragraph()

    def handle_data(self, data):
        text = data.replace('\r', '')
        if not text.strip() and len(self.current_runs) == 0:
            return
        fmt = self.format_stack[-1]
        self.current_runs.append({
            'text': text,
            'bold': fmt['bold'],
            'italic': fmt['italic'],
            'underline': fmt['underline'],
            'color': fmt['color'],
            'highlight': fmt['highlight']
        })

# Test sample HTML
sample_html = """
<h1 style="color: #1F4E79; text-align: center;">خطاب رسمي موحد</h1>
<p style="text-align: right;">هذا نص عادي مع كلمة <span style="color: rgb(220, 38, 38);">حمراء باللون المميز</span> وكلمة <span style="color: #16a34a; font-weight: bold;">خضراء وعريضة</span>.</p>
<p>أيضاً كلمة <span style="background-color: #fef08a;">مظللة باللون الأصفر</span> وكلمة <font color="#2563eb">زرقاء عبر وسم font</font>.</p>
<ul>
    <li>عنصر أول مع <span style="color: #7c3aed;">لون بنفسجي</span></li>
    <li>عنصر ثانٍ باللون الأسود المعتاد</li>
</ul>
"""

parser = EnhancedDocxHTMLParser()
parser.feed(sample_html)
parser._flush_paragraph()

print("Parsed Paragraphs count:", len(parser.paragraphs))
for i, p in enumerate(parser.paragraphs):
    print(f"P {i+1} (align={p['alignment']}, heading={p['heading_level']}):")
    for r in p['runs']:
        print(f"   Run: '{r['text']}' | bold={r['bold']} | color={r['color']} | highlight={r['highlight']}")

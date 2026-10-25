import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import copy

doc = docx.Document()
p = doc.add_paragraph()

r1 = p.add_run("Normal text. ")

r2 = p.add_run("Red text. ")
r2.font.color.rgb = RGBColor(220, 38, 38)

r3 = p.add_run("Navy text. ")
r3.font.color.rgb = RGBColor(31, 78, 121)

r4 = p.add_run("Green bold text. ")
r4.bold = True
r4.font.color.rgb = RGBColor(22, 163, 74)

# Highlight test
r5 = p.add_run("Highlighted yellow text. ")
rPr = r5._element.get_or_add_rPr()
highlight = OxmlElement('w:highlight')
highlight.set(qn('w:val'), 'yellow')
rPr.append(highlight)

doc.save("/Users/ahmadkhaled/Documents/GitHub/mixed/scratch/test_color.docx")
print("Saved successfully!")
